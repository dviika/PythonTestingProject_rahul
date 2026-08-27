import pandas as pd
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def compare_columns(input_file, output_file="comparison_result.xlsx"):
    df = pd.read_excel(input_file, dtype=str)

    if "OLD" not in df.columns or "NEW" not in df.columns:
        raise ValueError("Input file must contain columns named 'OLD' and 'NEW'")

    # Normalize: strip whitespace, lowercase for comparison
    old_series = df["OLD"].dropna().str.strip().str.lower()
    new_series = df["NEW"].dropna().str.strip().str.lower()

    old_set = set(old_series)compare
    new_set = set(new_series)

    only_in_old = sorted(old_set - new_set)
    only_in_new = sorted(new_set - old_set)
    in_both     = sorted(old_set & new_set)

    max_len = max(len(only_in_old), len(only_in_new), len(in_both))

    def pad(lst):
        return lst + [""] * (max_len - len(lst))

    result_df = pd.DataFrame({
        "Only in OLD (not in NEW)": pad(only_in_old),
        "Only in NEW (not in OLD)": pad(only_in_new),
        "In Both OLD and NEW":      pad(in_both),
    })

    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    header_fill   = PatternFill("solid", start_color="4472C4")
    old_only_fill = PatternFill("solid", start_color="FFB3B3")
    new_only_fill = PatternFill("solid", start_color="B3FFB3")
    both_fill     = PatternFill("solid", start_color="B3D9FF")
    col_fills     = [old_only_fill, new_only_fill, both_fill]

    for col_idx, col_name in enumerate(result_df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font      = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 30

    for row_idx, row in enumerate(result_df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = Font(name="Arial")
            cell.fill      = col_fills[col_idx - 1]
            cell.alignment = Alignment(horizontal="left")

    ws.row_dimensions[1].height = 30

    # Summary row
    summary_row = max_len + 3
    ws.cell(row=summary_row, column=1, value="Summary").font = Font(bold=True, name="Arial")
    ws.cell(row=summary_row + 1, column=1, value=f"Only in OLD: {len(only_in_old)} items")
    ws.cell(row=summary_row + 2, column=1, value=f"Only in NEW: {len(only_in_new)} items")
    ws.cell(row=summary_row + 3, column=1, value=f"In both: {len(in_both)} items")

    wb.save(output_file)
    print(f"Done! Results saved to: {output_file}")
    print(f"  Only in OLD : {len(only_in_old)}")
    print(f"  Only in NEW : {len(only_in_new)}")
    print(f"  In both     : {len(in_both)}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python compare_columns.py <input.xlsx> [output.xlsx]")
        sys.exit(1)

    input_path  = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else "comparison_result.xlsx"
    compare_columns(input_path, output_path)